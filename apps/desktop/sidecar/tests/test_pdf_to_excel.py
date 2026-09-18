"""Tests for `pdf.to_excel` -- extract tables from a PDF into an .xlsx workbook."""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

from ops import pdf_to_excel
from ops._common import OpError, noop_progress


def _make_table_pdf(path: Path, tables_per_page: list[list[list[list[str]]]]) -> None:
    """`tables_per_page` is one entry per page; each entry is a list of tables
    (each table a list of rows, each row a list of cell strings) to draw on
    that page. A page with an empty list gets a page of plain text instead --
    reportlab's SimpleDocTemplate can't easily emit a truly blank page, and a
    page pdfplumber can find zero tables on is exactly what these tests need."""
    from reportlab.platypus import PageBreak, Paragraph
    from reportlab.lib.styles import getSampleStyleSheet

    doc = SimpleDocTemplate(str(path), pagesize=letter)
    styles = getSampleStyleSheet()
    story = []
    for page_index, tables in enumerate(tables_per_page):
        if page_index > 0:
            story.append(PageBreak())
        if not tables:
            story.append(Paragraph("No table on this page, just text.", styles["Normal"]))
            continue
        for table_data in tables:
            t = Table(table_data)
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 1, colors.black)]))
            story.append(t)
    doc.build(story)


def test_to_excel_rejects_a_pdf_with_no_tables_anywhere(tmp_path, out_dir):
    src = tmp_path / "no_tables.pdf"
    _make_table_pdf(src, [[], []])  # two pages, neither has a table
    dest = out_dir / "o.xlsx"

    with pytest.raises(OpError) as exc:
        pdf_to_excel.run({"input": str(src), "output": str(dest)}, noop_progress)
    assert exc.value.code == "NO_TABLES_FOUND"
    assert not dest.exists()


def test_to_excel_writes_one_sheet_per_table_bearing_page(tmp_path, out_dir):
    src = tmp_path / "mixed.pdf"
    table_a = [["Name", "Age"], ["Alice", "30"]]
    table_c = [["City", "Pop"], ["Tokyo", "37M"]]
    _make_table_pdf(src, [[table_a], [], [table_c]])  # page 2 has no table
    dest = out_dir / "o.xlsx"

    result = pdf_to_excel.run({"input": str(src), "output": str(dest)}, noop_progress)

    assert result["sheets"] == 2
    assert result["output"] == str(dest)
    assert result["bytes"] > 0
    wb = openpyxl.load_workbook(str(dest))
    assert wb.sheetnames == ["Page 1", "Page 3"]  # page 2 contributes nothing
    rows_1 = list(wb["Page 1"].iter_rows(values_only=True))
    assert rows_1 == [("Name", "Age"), ("Alice", "30")]
    rows_3 = list(wb["Page 3"].iter_rows(values_only=True))
    assert rows_3 == [("City", "Pop"), ("Tokyo", "37M")]


def test_to_excel_progress_reaches_100(tmp_path, out_dir):
    src = tmp_path / "one.pdf"
    _make_table_pdf(src, [[[["A"], ["1"]]]])
    dest = out_dir / "o.xlsx"
    calls: list[int] = []

    pdf_to_excel.run({"input": str(src), "output": str(dest)}, lambda pct, note="": calls.append(pct))
    assert calls[-1] == 100
